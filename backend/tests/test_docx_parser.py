import tempfile
from pathlib import Path
import docx
from backend.app.ingestion.parser import DocumentParser


def test_docx_parsing():
    # Create temporary docx
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = Path(tmpdir) / "test_doc.docx"
        doc = docx.Document()
        doc.add_heading("TruthGuard Technical Specification", level=1)
        doc.add_paragraph("TruthGuard is an explainable NLP framework for hallucination detection.")
        doc.add_paragraph("Evidence is retrieved via hybrid BM25 and dense FAISS.")
        doc.save(str(docx_path))

        pages, total_pages = DocumentParser.parse_file(str(docx_path), "test_doc.docx")
        assert total_pages >= 1
        assert len(pages) == total_pages
        combined_text = " ".join(p["text"] for p in pages)
        assert "Technical Specification" in combined_text
        assert "hallucination detection" in combined_text
        assert "hybrid BM25" in combined_text
