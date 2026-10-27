import docx

doc = docx.Document(r"docs\ADA23DLU04_PROJECT_REPORT_WITH_FIGURES.docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for idx, p in enumerate(doc.paragraphs):
    text = p.text
    align = str(p.alignment)
    line_spacing = p.paragraph_format.line_spacing
    sb = p.paragraph_format.space_before
    sa = p.paragraph_format.space_after
    li = p.paragraph_format.left_indent
    fi = p.paragraph_format.first_line_indent
    drawings = len(p._element.xpath(".//w:drawing"))
    
    # Check for formatting red flags
    flags = []
    if drawings > 0:
        flags.append(f"DRAWING({drawings})")
    if len(text.strip()) == 0 and drawings == 0:
        flags.append("BLANK_PARAGRAPH")
    if "\n" in text:
        flags.append(f"MANUAL_NEWLINE({text.count(chr(10))})")
    if "  " in text:
        flags.append(f"DOUBLE_SPACES({text.count('  ')})")
    if "•" in text:
        flags.append("BULLET_CHAR")
    if any(ord(c) > 127 for c in text):
        non_ascii = [f"U+{ord(c):04X}({c})" for c in text if ord(c) > 127]
        flags.append(f"NON_ASCII({','.join(set(non_ascii))})")

    # Flag summary
    flag_str = " | ".join(flags) if flags else "OK"
    preview = text.strip()[:50].replace("\n", " [NL] ")
    print(f"P{idx:03d} | align={align:12s} | sp_b={str(sb):6s} | sp_a={str(sa):6s} | flags={flag_str}")
    if flags and "OK" not in flags:
        print(f"       -> '{preview}'")
